from __future__ import annotations

import html
import io
import re
import unittest
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from shared.config import ConfigError
from sponsor_deadline_mails.core import generate_deadline_mails
from sponsor_deadline_mails.parser import resolve_workbook_layout


# Match the new workbook's headers and column positions, with synthetic contacts.
MUC26_HEADERS = (
    " ", "Sponsor", "AP", "Deal liegt vor", "Paket ", "Gesprächsanzahl",
    "Sponsorinfos erhalten", "Upgrade 1", "Upgrade 2", "Standnummer", "Sprache",
    "ASP 1 Vorname", "ASP 1 Nachname", "ASP 1 E-Mail Adresse",
    "ASP 2 Vorname", "ASP 2 Nachname", "ASP 2 E-Mail Adresse", "ASP 3 Vorname",
    "Logo", "Sponsoren Vor Ort", "Stream 1 & 2", "Handout",
    "Booklet Informationen", "Vortragsinformationen", "Onboarding",
    "Target Account Liste erhalten", "LED Wand Design", "LED Wand besonderheiten",
    "LED Wand Designs Anzahl", "Posting erhalten", "Posting gepostet",
    "Präsentation erhalten", "Teilnehmerliste Gesprächswünsche erhalten", "Zusatzinfos ",
)
CHECK_HEADERS = (
    "Logo", "Sponsoren Vor Ort", "Handout", "Booklet Informationen",
    "Vortragsinformationen", "Onboarding", "Target Account Liste erhalten",
    "LED Wand Design", "Posting gepostet", "Präsentation erhalten",
    "Teilnehmerliste Gesprächswünsche erhalten",
)


def muc26_workbook_bytes(*, updates=None, reordered=False, header_row=1):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Deals"
    headers = list(MUC26_HEADERS)
    if reordered:
        headers = ["Eventnotiz", *reversed(headers)]
    if header_row > 1:
        worksheet["A1"] = "Synthetic sponsor preparation"
    values = {
        "Sponsor": "Synthetic MUC Sponsor",
        "Deal liegt vor": "Check",
        "Paket ": "Gold",
        "Sponsorinfos erhalten": "Check",
        "Sprache": "DE",
        "ASP 1 Vorname": "Alex",
        "ASP 1 Nachname": "Example",
        "ASP 1 E-Mail Adresse": "primary@example.com",
        "ASP 2 Vorname": "Robin",
        "ASP 2 Nachname": "Sample",
        "ASP 2 E-Mail Adresse": "secondary@example.com",
        "ASP 3 Vorname": "ThirdContactWithoutEmail",
        "Posting erhalten": "Check",
        "Zusatzinfos ": "Internal note, not a status",
        **{header: "Check" for header in CHECK_HEADERS},
        **(updates or {}),
    }
    for column, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column, header)
        worksheet.cell(header_row + 1, column, values.get(header))
    workbook.create_sheet("Bookletinformationen").append(["Title", "Description"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def task_row(body, task_fragment):
    """Locate a task and its status in the generated customer-facing HTML."""
    matches = [
        row for row in re.findall(r"<tr\b[^>]*>.*?</tr>", body, flags=re.DOTALL)
        if task_fragment in html.unescape(row)
    ]
    if len(matches) != 1:
        raise AssertionError(f"Expected one task row for {task_fragment!r}, found {len(matches)}")
    return matches[0]


class SponsorDeadlineMailsMuc26Tests(unittest.TestCase):
    def generate(self, updates=None, *, today=date(2026, 10, 1), **kwargs):
        return generate_deadline_mails(
            muc26_workbook_bytes(updates=updates), today=today, **kwargs
        ).mails[0]

    def assert_task_status(self, mail, fragment, expected):
        labels = {"green": "Bereits erhalten", "red": "Handlungsbedarf", "yellow": "Ausstehend"}
        self.assertIn(labels[expected], task_row(mail.html_body, fragment))

    def test_new_layout_maps_shifted_status_columns_and_ignores_unrelated_headers(self):
        workbook = load_workbook(io.BytesIO(muc26_workbook_bytes()))
        self.addCleanup(workbook.close)

        layout = resolve_workbook_layout(workbook["Deals"], strict=True)

        self.assertEqual(layout.columns["sponsor_name"], "B")
        self.assertEqual(layout.columns["target_accounts"], "Z")
        self.assertEqual(layout.columns["posting_published"], "AE")
        self.assertEqual(layout.columns["presentation"], "AF")
        self.assertEqual(layout.columns["meeting_selection"], "AG")
        self.assertTrue({"G", "R", "AD", "AH"}.isdisjoint(layout.columns.values()))

    def test_new_workbook_generates_munich_mail_without_false_pending_tasks(self):
        mail = self.generate()

        self.assertEqual(mail.sponsor_name, "Synthetic MUC Sponsor")
        self.assertEqual((mail.to_email, mail.cc_email), ("primary@example.com", "secondary@example.com"))
        self.assertEqual((mail.green_count, mail.red_count, mail.yellow_count, mail.white_count), (11, 0, 1, 2))
        self.assertIn("17.11.2026", mail.html_body)
        self.assertIn("19.11.2026", mail.html_body)
        self.assert_task_status(mail, "Sende deine Target Account Liste", "green")

    def test_received_visual_does_not_complete_published_post_from_shifted_presentation(self):
        # AD is received, AE is published, and AF is presentation in MUC26.
        mail = self.generate({"Posting erhalten": "Check", "Posting gepostet": None})

        self.assertEqual((mail.green_count, mail.red_count), (10, 1))
        self.assert_task_status(mail, "Poste das individuelle Visual", "red")
        self.assert_task_status(mail, "Vortragspräsentation.", "green")
        self.assert_task_status(mail, "Sende Severin die Auswahl", "green")

    def test_meeting_selection_check_cannot_complete_missing_presentation(self):
        mail = self.generate({"Präsentation erhalten": None})

        self.assertEqual((mail.green_count, mail.red_count), (10, 1))
        self.assert_task_status(mail, "Poste das individuelle Visual", "green")
        self.assert_task_status(mail, "Vortragspräsentation.", "red")
        self.assert_task_status(mail, "Sende Severin die Auswahl", "green")

    def test_extra_notes_cannot_complete_missing_meeting_selection(self):
        updates = {"Teilnehmerliste Gesprächswünsche erhalten": None, "Zusatzinfos ": "Check"}
        for today, status in ((date(2026, 10, 1), "yellow"), (date(2026, 11, 10), "red")):
            with self.subTest(today=today):
                mail = self.generate(updates, today=today)
                self.assert_task_status(mail, "Sende Severin die Auswahl", status)
                self.assert_task_status(mail, "Vortragspräsentation.", "green")
                self.assertEqual(mail.green_count, 10)

    def test_checked_meeting_selection_is_independent_of_extra_notes(self):
        for note in (None, "Check", "Waiting for information"):
            with self.subTest(note=note):
                mail = self.generate({"Zusatzinfos ": note})
                self.assert_task_status(mail, "Sende Severin die Auswahl", "green")
                self.assertEqual(mail.red_count, 0)

    def test_general_sponsor_information_check_does_not_complete_open_materials(self):
        mail = self.generate({
            "Sponsorinfos erhalten": "Check", "Logo": None, "Handout": None,
            "Booklet Informationen": None, "Vortragsinformationen": None,
            "LED Wand Design": None,
        })

        self.assertEqual((mail.green_count, mail.red_count), (6, 5))
        for task in (
            "Sende das Unternehmenslogo", "Sende das Handout/Whitepaper.",
            "Sende die Informationen für das Booklet.", "Vortragsinformationen.",
            "Sende das LED-Wand-Design.",
        ):
            with self.subTest(task=task):
                self.assert_task_status(mail, task, "red")

    def test_reordered_new_headers_keep_contacts_and_opposing_statuses(self):
        updates = {"Posting gepostet": None, "Zusatzinfos ": "Check"}
        baseline = self.generate(updates)

        shifted = generate_deadline_mails(
            muc26_workbook_bytes(updates=updates, reordered=True, header_row=3),
            today=date(2026, 10, 1),
        ).mails[0]

        self.assertEqual(shifted.row_number, 4)
        self.assertEqual((shifted.to_email, shifted.cc_email), (baseline.to_email, baseline.cc_email))
        self.assertEqual(shifted.html_body, baseline.html_body)

    def test_english_contact_fallback_ignores_third_contact_without_email_for_next_event(self):
        mail = self.generate({
            "Sprache": "ENG", "ASP 1 E-Mail Adresse": None,
        }, event_city="Hamburg", event_start=date(2027, 2, 16), event_end=date(2027, 2, 18))

        self.assertEqual(mail.language, "EN")
        self.assertEqual((mail.to_email, mail.cc_email), ("secondary@example.com", ""))
        self.assertIn("Hi Robin,", mail.html_body)
        self.assertIn("Hamburg", mail.html_body)
        self.assertIn("16/02/2027", mail.html_body)
        self.assertIn("14/01/2027", mail.html_body)
        self.assertIn("28/01/2027", mail.html_body)
        self.assertNotIn("ThirdContactWithoutEmail", mail.html_body)
        self.assertNotIn("Allianz Arena", mail.html_body)
        self.assertNotIn("mysecurityevent-munich", mail.html_body)
        self.assertEqual((mail.green_count, mail.red_count), (11, 0))

    def test_page_upload_generates_new_workbook_preview_without_sending(self):
        excel_bytes = muc26_workbook_bytes(updates={"Posting gepostet": None})
        upload = SimpleNamespace(name="00_MUC26_Master_SPO_Infos.xlsx", getvalue=lambda: excel_bytes)
        with (
            patch("streamlit.file_uploader", return_value=upload),
            patch("shared.config.load_imap_draft_settings", side_effect=ConfigError("Test IMAP disabled")),
            patch("shared.config.load_smtp_send_settings", side_effect=ConfigError("Test SMTP disabled")),
            patch("sponsor_deadline_mails.create_smtp_sends") as smtp,
            patch("sponsor_deadline_mails.create_imap_drafts") as imap,
        ):
            app = AppTest.from_file("pages/Sponsoren_Deadline_Mails.py", default_timeout=20).run()
            self.assertEqual([], [exception.message for exception in app.exception])

            next(button for button in app.button if button.label == "Generieren").click().run()

            self.assertEqual([], [exception.message for exception in app.exception])
            result = app.session_state["sdm_result"]
            self.assertEqual(result.processed_count, 1)
            self.assertEqual(result.mails[0].sponsor_name, "Synthetic MUC Sponsor")
            self.assertEqual(result.mails[0].red_count, 1)
            self.assertEqual(len(app.dataframe), 1)
            smtp.assert_not_called()
            imap.assert_not_called()


if __name__ == "__main__":
    unittest.main()
