from __future__ import annotations

import io
import unittest
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree

from openpyxl import Workbook

from sponsor_deadline_mails.core import generate_deadline_mails, parse_event_date


HEADERS = [
    "Nr", "Name", "AP", "Deal liegt vor", "Paket", "Gesprächsanzahl",
    "Upgrade 1", "Upgrade 2", "Upgrade 3", "Standnummer", "Sprache",
    "ASP 1 Vorname", "ASP 1 Nachname", "ASP 1 E-Mail Adresse",
    "ASP 2 Vorname", "ASP 2 Nachname", "ASP 2 E-Mail Adresse", "ASP 2 Telefonnummer",
    "Logo", "Sponsoren Vor Ort", "Stream 1 & und 2", "Handout",
    "Booklet Informationen", "Vortragsinformationen", "Onboarding",
    "Wunschteilnehmerliste erhalten", "LED Wand Design", "LED Wand besonderheiten",
    "LED Wand Anzahl", "Pitch Night Thema", "Posting erhalten", "Posting gepostet",
    "Präsentation erhalten", "Teilnehmerliste Gesprächswünsche erhalten",
]
CHECK_HEADERS = [
    "Logo", "Sponsoren Vor Ort", "Handout", "Booklet Informationen", "Vortragsinformationen",
    "Onboarding", "Wunschteilnehmerliste erhalten", "LED Wand Design", "Posting gepostet",
    "Präsentation erhalten", "Teilnehmerliste Gesprächswünsche erhalten",
]


def workbook_bytes(*, shifted=False, updates=None, missing=None, header_row=1):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Deals"
    headers = [header for header in HEADERS if header != missing]
    if shifted:
        headers = ["Zusätzliche Eventnotiz", *reversed(headers)]
    if header_row > 1:
        worksheet["A1"] = "Sponsor preparation"
    for column, header in enumerate(headers, 1):
        worksheet.cell(header_row, column, header)
    values = {
        "Name": "Example Sponsor", "Deal liegt vor": "Check", "Paket": "Gold",
        "Sprache": "DE", "ASP 1 Vorname": "Alex", "ASP 1 E-Mail Adresse": "alex@example.com",
        "ASP 2 E-Mail Adresse": "copy@example.com",
        **{header: "Check" for header in CHECK_HEADERS},
        **(updates or {}),
    }
    for column, header in enumerate(headers, 1):
        worksheet.cell(header_row + 1, column, values.get(header))
    workbook.create_sheet("Bookletinformationen").append(["Text", "Titel"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def with_formula_cache(
    excel_bytes: bytes,
    coordinate: str,
    cached_value: str,
    *,
    cell_type: str | None = None,
) -> bytes:
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(excel_bytes)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            contents = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                tree = ElementTree.fromstring(contents)
                cell = tree.find(f'.//s:c[@r="{coordinate}"]', namespace)
                if cell_type is not None:
                    cell.set("t", cell_type)
                cell.find("s:v", namespace).text = cached_value
                contents = ElementTree.tostring(tree, encoding="utf-8")
            target.writestr(entry, contents)
    return output.getvalue()


class SponsorDeadlineMailsCoreTests(unittest.TestCase):
    def generate(self, excel_bytes=None, **kwargs):
        return generate_deadline_mails(
            excel_bytes or workbook_bytes(), today=date(2026, 10, 1), **kwargs
        )

    def test_checked_munich_sponsor_has_no_false_action_required(self):
        result = self.generate()
        self.assertEqual((result.event_city, result.event_start, result.event_end), (
            "München", date(2026, 11, 17), date(2026, 11, 19)
        ))
        mail = result.mails[0]
        self.assertEqual((mail.green_count, mail.red_count, mail.yellow_count, mail.white_count), (11, 0, 1, 2))
        for text in ("18.08.2026", "31.08.2026", "15.10.2026", "29.10.2026", "03.11.2026", "10.11.2026", "17.11.2026"):
            self.assertIn(text, mail.html_body)
        self.assertIn('href="https://calendly.com/mysecurityevent/1-1-meeting-sponsoren-mysecurityevent-munich"', mail.html_body)
        self.assertIn("Business Club der Allianz Arena", mail.html_body)
        self.assertIn("14:00", mail.html_body)
        self.assertIn("16:00", mail.html_body)
        self.assertNotIn("26.03.2026", mail.html_body)

    def test_reordered_columns_and_lower_header_preserve_contacts_and_statuses(self):
        baseline = self.generate().mails[0]
        shifted = self.generate(workbook_bytes(shifted=True, header_row=3)).mails[0]
        self.assertEqual(shifted.row_number, 4)
        self.assertEqual((shifted.to_email, shifted.cc_email), ("alex@example.com", "copy@example.com"))
        self.assertEqual(shifted.html_body, baseline.html_body)
        self.assertEqual(shifted.red_count, 0)

    def test_english_followup_event_has_new_dates_and_no_old_event_details(self):
        mail = self.generate(
            workbook_bytes(updates={"Sprache": "ENG"}), event_city="Hamburg",
            event_start=date(2027, 2, 16), event_end=date(2027, 2, 18),
        ).mails[0]
        for text in ("17/11/2026", "30/11/2026", "14/01/2027", "28/01/2027", "02/02/2027", "09/02/2027", "16/02/2027"):
            self.assertIn(text, mail.html_body)
        self.assertNotIn("Allianz Arena", mail.html_body)
        self.assertNotIn("mysecurityevent-munich", mail.html_body)
        self.assertIn("Hamburg", mail.html_body)

    def test_same_date_different_city_does_not_inherit_munich_links(self):
        mail = self.generate(event_city="Hamburg").mails[0]
        self.assertNotIn("Allianz Arena", mail.html_body)
        self.assertNotIn("mysecurityevent-munich", mail.html_body)

    def test_post_received_does_not_replace_published_check(self):
        mail = self.generate(workbook_bytes(updates={"Posting gepostet": None, "Posting erhalten": "Check"})).mails[0]
        self.assertEqual(mail.red_count, 1)

    def test_cached_formula_result_is_used_instead_of_formula_text(self):
        excel_bytes = workbook_bytes(updates={"Logo": "=FALSE()"})
        mail = self.generate(with_formula_cache(excel_bytes, "S2", "0")).mails[0]
        self.assertEqual(mail.red_count, 1)
        self.assertEqual(mail.green_count, 10)

    def test_boolean_formula_caches_preserve_completed_and_pending_statuses(self):
        excel_bytes = workbook_bytes(updates={"Logo": "=TRUE()"})
        for cached_value, expected_red in (("1", 0), ("0", 1)):
            with self.subTest(cached_value=cached_value):
                mail = self.generate(with_formula_cache(
                    excel_bytes, "S2", cached_value, cell_type="b"
                )).mails[0]
                self.assertEqual(mail.red_count, expected_red)
                self.assertEqual(mail.green_count, 11 - expected_red)

    def test_intentionally_empty_string_formula_cache_means_pending(self):
        excel_bytes = workbook_bytes(updates={"Logo": '=IF(FALSE(),"Check","")'})

        mail = self.generate(with_formula_cache(excel_bytes, "S2", "", cell_type="str")).mails[0]

        self.assertEqual(mail.red_count, 1)
        self.assertEqual(mail.green_count, 10)

    def test_intentionally_empty_sponsor_name_formula_cache_skips_row(self):
        excel_bytes = workbook_bytes(updates={"Name": '=IF(FALSE(),"Example Sponsor","")'})

        result = self.generate(with_formula_cache(excel_bytes, "B2", "", cell_type="str"))

        self.assertEqual(result.processed_count, 0)
        self.assertEqual(result.skipped_count, 0)

    def test_inactive_sponsor_does_not_require_unused_formula_caches(self):
        result = self.generate(workbook_bytes(updates={
            "Deal liegt vor": "no", "Logo": "=TRUE()",
            "ASP 1 E-Mail Adresse": '=CONCAT("alex", "@example.com")',
        }))

        self.assertEqual(result.processed_count, 0)
        self.assertEqual(result.skipped_count, 1)

    def test_sponsor_without_contact_does_not_require_unused_formula_caches(self):
        result = self.generate(workbook_bytes(updates={
            "ASP 1 E-Mail Adresse": None, "ASP 2 E-Mail Adresse": None,
            "Logo": "=TRUE()", "Paket": '=CONCAT("G", "old")',
        }))

        self.assertEqual(result.processed_count, 0)
        self.assertEqual(result.skipped_count, 1)

    def test_non_talk_package_does_not_require_unused_talk_formula_caches(self):
        mail = self.generate(workbook_bytes(updates={
            "Paket": "Premium", "Vortragsinformationen": "=TRUE()",
            "Präsentation erhalten": "=TRUE()",
        })).mails[0]

        self.assertEqual(mail.green_count, 9)
        self.assertEqual(mail.red_count, 0)
        self.assertNotIn("Vortragsinformationen.", mail.html_body)
        self.assertNotIn("Vortragspräsentation.", mail.html_body)

    def test_unused_second_contact_names_do_not_require_formula_caches(self):
        mail = self.generate(workbook_bytes(updates={
            "ASP 2 Vorname": '=CONCAT("A", "lex")',
            "ASP 2 Nachname": '=CONCAT("Ex", "ample")',
        })).mails[0]

        self.assertEqual(mail.to_email, "alex@example.com")
        self.assertIn("Hallo Alex,", mail.html_body)

    def test_contact_formula_needed_to_select_recipient_requires_cache(self):
        with self.assertRaisesRegex(ValueError, "Deals!N2.*kein gespeichertes Ergebnis"):
            self.generate(workbook_bytes(updates={
                "ASP 1 E-Mail Adresse": '=CONCAT("alex", "@example.com")',
            }))

    def test_chosen_second_contact_name_requires_formula_cache(self):
        with self.assertRaisesRegex(ValueError, "Deals!O2.*kein gespeichertes Ergebnis"):
            self.generate(workbook_bytes(updates={
                "ASP 1 E-Mail Adresse": None,
                "ASP 2 Vorname": '=CONCAT("A", "lex")',
            }))

    def test_uncached_formula_is_reported_with_cell_coordinate(self):
        with self.assertRaisesRegex(ValueError, "Deals!S2.*kein gespeichertes Ergebnis"):
            self.generate(workbook_bytes(updates={"Logo": "=TRUE()"}))

    def test_missing_task_column_and_wrong_sheet_are_reported(self):
        with self.assertRaisesRegex(ValueError, "Handout"):
            self.generate(workbook_bytes(missing="Handout"))
        with self.assertRaisesRegex(ValueError, "Kopfzeile|Pflichtspalten"):
            self.generate(sheet_name="Bookletinformationen")

    def test_event_end_before_start_is_rejected_by_engine(self):
        with self.assertRaisesRegex(ValueError, "Event-Ende"):
            self.generate(event_end=date(2026, 11, 16))

    def test_invalid_onboarding_link_and_checkin_times_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "Onboarding-Link"):
            self.generate(onboarding_url="javascript:alert(1)")
        with self.assertRaisesRegex(ValueError, "Check-ins"):
            self.generate(checkin_start="16:00", checkin_end="14:00")

    def test_formula_sponsor_name_without_cache_does_not_silently_skip_row(self):
        with self.assertRaisesRegex(ValueError, "Deals!B2"):
            self.generate(workbook_bytes(updates={"Name": '=CONCAT("Example", " Sponsor")'}))

    def test_excel_error_is_reported_instead_of_interpreted_as_status(self):
        with self.assertRaisesRegex(ValueError, "Excel-Fehler in Deals!S2"):
            self.generate(workbook_bytes(updates={"Logo": "#VALUE!"}))

    def test_datetime_input_is_normalized_to_date(self):
        self.assertEqual(parse_event_date(datetime(2026, 11, 17, 14, 30)), date(2026, 11, 17))


if __name__ == "__main__":
    unittest.main()
