from __future__ import annotations

import io
import unittest
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from sponsor_deadline_mails.core import generate_deadline_mails
from sponsor_deadline_mails.parser import resolve_workbook_layout
from sponsor_deadline_mails.workbook_statuses import build_supplemental_received


DEALS_HEADERS = [
    "Sponsor", "Paket", "Deal liegt vor", "ASP 1 E-Mail Adresse", "Logo",
    "Sponsoren Vor Ort", "Handout", "Booklet Informationen", "Vortragsinformationen",
    "Onboarding", "Target Account Liste erhalten", "LED Wand Design", "Posting gepostet",
    "Präsentation erhalten", "Teilnehmerliste Gesprächswünsche erhalten",
]
TALK_HEADERS = [
    "Sponsor", "Logo", "Vortragstitel", "Vortrag Bullet Points",
    "Vortrag Speaker Vorname", "Vortrag Speaker Nachname", "Vortrag Speaker Jobtitel",
    "Foto Vortragsspeaker", "Vortragspräsentation erhalten",
    "Vortragsinformation in der Druckversion eingetragen",
]
TALK_VALUES = [
    "Acme GmbH", "Check", "A security talk", "Useful points", "Alex", "Example",
    "Engineer", "Check", "Check", None,
]


def make_workbook(*, package="Gold", duplicate_deals=False):
    workbook = Workbook()
    deals = workbook.active
    deals.title = "Deals"
    deals.append(DEALS_HEADERS)
    values = ["Acme GmbH", package, "Check", "contact@example.com"]
    deals.append(values)
    if duplicate_deals:
        deals.append(values)
    return workbook


def workbook_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def with_cache(excel_bytes, sheet, coordinate, value, *, cell_type="str"):
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(excel_bytes)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            contents = source.read(entry.filename)
            if entry.filename == f"xl/worksheets/sheet{sheet}.xml":
                tree = ElementTree.fromstring(contents)
                cell = tree.find(f'.//s:c[@r="{coordinate}"]', namespace)
                cell.set("t", cell_type)
                cell.find("s:v", namespace).text = value
                contents = ElementTree.tostring(tree, encoding="utf-8")
            target.writestr(entry, contents)
    return output.getvalue()


class WorkbookStatusTests(unittest.TestCase):
    def setUp(self):
        self.workbook = make_workbook()
        self.addCleanup(self.workbook.close)

    def add_talk_sheet(self, *, values=None, title="Vortragsinformationen"):
        sheet = self.workbook.create_sheet(title)
        sheet.append(TALK_HEADERS)
        sheet.append(TALK_VALUES if values is None else values)
        return sheet

    def resolve(self, excel_bytes=None, **kwargs):
        excel_bytes = excel_bytes or workbook_bytes(self.workbook)
        source = load_workbook(io.BytesIO(excel_bytes), data_only=False)
        values = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        self.addCleanup(source.close)
        self.addCleanup(values.close)
        return build_supplemental_received(
            source, values, sheet_name="Deals", layout=resolve_workbook_layout(source["Deals"], strict=True),
            **kwargs,
        )

    def generate(self, excel_bytes=None, **kwargs):
        return generate_deadline_mails(
            excel_bytes or workbook_bytes(self.workbook), today=date(2026, 10, 1), **kwargs,
        )

    def test_companion_checks_and_complete_talk_add_completion(self):
        booklet = self.workbook.create_sheet("Bookletinformationen")
        booklet.append(["Sponsor", "Bookletinformationen liegen vor"])
        booklet.append(["Acme GmbH", "Check"])
        self.add_talk_sheet()

        self.assertEqual(self.resolve(), {2: frozenset({"booklet", "logo", "talk_info", "presentation"})})
        mail = self.generate().mails[0]
        self.assertEqual(mail.green_count, 4)
        self.assertEqual(mail.red_count, 6)

    def test_blank_companion_cells_never_reset_main_checks(self):
        deals = self.workbook["Deals"]
        for key in ["Logo", "Booklet Informationen", "Vortragsinformationen", "Präsentation erhalten"]:
            deals.cell(2, DEALS_HEADERS.index(key) + 1, "Check")
        self.add_talk_sheet(values=["Acme GmbH"])

        self.assertEqual(self.resolve(), {})
        self.assertEqual(self.generate().mails[0].green_count, 4)

    def test_no_talk_tasks_for_current_premium_package(self):
        self.workbook["Deals"]["B2"] = "Premium"
        talk = self.add_talk_sheet()
        talk["I2"] = "=TRUE()"  # Irrelevant talk formula must not block Premium.

        self.assertEqual(self.resolve(), {2: frozenset({"logo"})})
        mail = self.generate().mails[0]
        self.assertEqual(mail.green_count, 1)
        self.assertNotIn("Vortragspräsentation.", mail.html_body)

    def test_reordering_and_lower_headers_work_for_followup_event(self):
        talk = self.workbook.create_sheet(" vortragsinformationen ")
        talk.append(["Next event details"])
        talk.append([])
        talk.append(list(reversed(TALK_HEADERS)))
        values = TALK_VALUES.copy()
        values[0] = "  ACME   GmbH  "
        talk.append(list(reversed(values)))

        self.assertEqual(self.resolve(), {2: frozenset({"logo", "talk_info", "presentation"})})
        result = self.generate(event_city="Hamburg", event_start=date(2027, 2, 16), event_end=date(2027, 2, 18))
        self.assertEqual(result.mails[0].green_count, 3)
        self.assertIn("28.01.2027", result.mails[0].html_body)

    def test_partial_talk_or_print_version_check_is_not_complete(self):
        values = TALK_VALUES.copy()
        values[5] = None
        values[9] = "Check"
        self.add_talk_sheet(values=values)

        self.assertEqual(self.resolve(), {2: frozenset({"logo", "presentation"})})

    def test_missing_photo_check_prevents_talk_completion(self):
        values = TALK_VALUES.copy()
        values[7] = None
        self.add_talk_sheet(values=values)

        self.assertNotIn("talk_info", self.resolve()[2])

    def test_placeholders_are_not_complete_talk_content(self):
        self.add_talk_sheet()
        for placeholder in ["-", "tbd", "TBA", "folgt", "offen", "...", "wird nachgereicht"]:
            with self.subTest(placeholder=placeholder):
                self.workbook["Vortragsinformationen"]["G2"] = placeholder
                self.assertNotIn("talk_info", self.resolve()[2])

    def test_archives_deals_copies_and_unrelated_names_are_ignored(self):
        self.add_talk_sheet(title="Vortragsinformationen Archiv")
        self.add_talk_sheet(title="Deals Kopie")
        values = TALK_VALUES.copy()
        values[0] = "Acme GmbH & Partner"
        self.add_talk_sheet(values=values)

        self.assertEqual(self.resolve(), {})

    def test_duplicate_deals_with_relevant_evidence_raise_clear_error(self):
        self.workbook["Deals"].append([" ACME GmbH ", "Gold", "Check", "other@example.com"])
        self.add_talk_sheet()

        with self.assertRaisesRegex(ValueError, "Mehrdeutige Sponsorzuordnung.*Deals.*2, 3"):
            self.resolve()

    def test_duplicate_companion_names_with_evidence_raise_clear_error(self):
        talk = self.add_talk_sheet()
        talk.append(TALK_VALUES)

        with self.assertRaisesRegex(ValueError, "Mehrdeutige Sponsorzuordnung.*Vortragsinformationen.*2, 3"):
            self.resolve()

    def test_duplicates_without_relevant_evidence_do_not_block(self):
        self.workbook["Deals"].append(["Acme GmbH", "Gold", "Check", "other@example.com"])
        self.add_talk_sheet(values=["Acme GmbH"])

        self.assertEqual(self.resolve(), {})

    def test_companion_name_formula_uses_validated_cache(self):
        talk = self.add_talk_sheet()
        talk["A2"] = "=Deals!A2"
        excel_bytes = workbook_bytes(self.workbook)
        with self.assertRaisesRegex(ValueError, "Vortragsinformationen!A2.*kein gespeichertes Ergebnis"):
            self.resolve(excel_bytes)

        excel_bytes = with_cache(excel_bytes, 2, "A2", "Acme GmbH")
        self.assertEqual(self.resolve(excel_bytes), {2: frozenset({"logo", "talk_info", "presentation"})})

    def test_needed_uncached_status_and_content_formulas_raise(self):
        talk = self.add_talk_sheet()
        talk["I2"] = "=TRUE()"
        with self.assertRaisesRegex(ValueError, "Vortragsinformationen!I2.*kein gespeichertes Ergebnis"):
            self.resolve()
        talk["I2"] = "Check"
        talk["C2"] = '=CONCAT("Security", " talk")'
        with self.assertRaisesRegex(ValueError, "Vortragsinformationen!C2.*kein gespeichertes Ergebnis"):
            self.resolve()

    def test_obvious_incomplete_details_skip_irrelevant_content_formulas(self):
        talk = self.add_talk_sheet()
        talk["C2"] = '=CONCAT("Security", " talk")'
        talk["G2"] = "tbd"
        self.assertEqual(self.resolve(), {2: frozenset({"logo", "presentation"})})

    def test_unrelated_number_stand_and_archive_formulas_do_not_block(self):
        booklet = self.workbook.create_sheet("Bookletinformationen")
        booklet.append(["Nummer", "Sponsor", "Bookletinformationen liegen vor"])
        booklet.append(["=Deals!Z2", "Acme GmbH", "Check"])
        self.workbook.create_sheet("Standnummer").append(["=TRUE()"])
        self.workbook.create_sheet("Bookletinformationen Archiv").append(["=TRUE()"])

        self.assertEqual(self.resolve(), {2: frozenset({"booklet"})})

    def test_unneeded_secondary_status_formula_does_not_block_existing_check(self):
        self.workbook["Deals"]["N2"] = "Check"
        talk = self.add_talk_sheet()
        talk["I2"] = "=TRUE()"

        self.assertEqual(self.resolve(), {2: frozenset({"logo", "talk_info"})})
        self.assertEqual(self.generate().mails[0].green_count, 3)

    def test_companion_without_applicable_statuses_does_not_need_formula_cache(self):
        self.workbook["Deals"]["B2"] = "Premium"
        talk = self.workbook.create_sheet("Vortragsinformationen")
        talk.append(["Sponsor", "Vortragspräsentation erhalten"])
        talk.append(["=Deals!A2", "=TRUE()"])

        self.assertEqual(self.resolve(), {})

    def test_boolean_numeric_and_completion_markers_are_not_talk_text(self):
        talk = self.add_talk_sheet()
        for value in [False, True, 0, 1, 42, "Check", "true", "1", "✓", "received"]:
            with self.subTest(value=value):
                talk["G2"] = value
                self.assertNotIn("talk_info", self.resolve()[2])

    def test_old_or_unparseable_talk_date_rejects_talk_but_keeps_logo(self):
        talk = self.add_talk_sheet()
        talk["K1"] = "Vortragsdatum"
        for value in ["09.05.2024", date(2024, 5, 9), "tbd", True, 1]:
            with self.subTest(value=value):
                talk["K2"] = value
                self.assertEqual(self.resolve(), {2: frozenset({"logo"})})
                self.assertEqual(self.generate().mails[0].green_count, 1)

    def test_current_and_future_events_use_their_own_talk_date_range(self):
        talk = self.add_talk_sheet()
        talk["K1"] = "Vortragsdatum"
        for value in ["17.11.2026", "2026-11-19", datetime(2026, 11, 18, 12, 0)]:
            with self.subTest(value=value):
                talk["K2"] = value
                self.assertIn("presentation", self.resolve()[2])
        talk["K2"] = "17.02.2027"
        self.assertEqual(self.resolve(), {2: frozenset({"logo"})})
        future = dict(event_start=date(2027, 2, 16), event_end=date(2027, 2, 18))
        self.assertEqual(self.resolve(**future), {2: frozenset({"logo", "talk_info", "presentation"})})
        self.assertEqual(self.generate(event_city="Hamburg", **future).mails[0].green_count, 3)

    def test_needed_talk_date_formula_requires_cache_but_irrelevant_one_does_not(self):
        talk = self.add_talk_sheet()
        talk["K1"] = "Vortragsdatum"
        talk["K2"] = '=DATE(2026,11,17)'
        with self.assertRaisesRegex(ValueError, "Vortragsinformationen!K2.*kein gespeichertes Ergebnis"):
            self.resolve()
        cached = with_cache(workbook_bytes(self.workbook), 2, "K2", "17.11.2026")
        self.assertIn("talk_info", self.resolve(cached)[2])
        self.workbook["Deals"]["B2"] = "Premium"
        self.assertEqual(self.resolve(), {2: frozenset({"logo"})})


if __name__ == "__main__":
    unittest.main()
